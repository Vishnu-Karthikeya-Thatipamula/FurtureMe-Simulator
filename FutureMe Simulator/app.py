import streamlit as st
import pandas as pd
import plotly.express as px
import hashlib
import json
import os
import requests 
import mysql.connector 
from google import genai
import re
import tempfile
import time
import PyPDF2
import docx
from google.genai import errors as genai_errors
from google.genai import types as genai_types

# ==========================================
# DATABASE & API CONFIGURATION
# ==========================================
def get_db_connection():
    return mysql.connector.connect(
        host="localhost", user="root", password="Vishnu@134",
        database="futureme_db", port=3306, use_pure=True
    )

GOOGLE_CLIENT_ID = "Your_Google_Client_ID"
GOOGLE_CLIENT_SECRET = "Your_Google_Client_Secret"
REDIRECT_URI = "http://localhost:8501"

# --- NEW SDK CLIENT INITIALIZATION ---
# NOTE: Gemini model names are case-sensitive and must exist for your API key.
# Verify with: [m.name for m in gemini_client.models.list()]
GEMINI_API_KEY = "Your_Gemini_API_Key"
GEMINI_INTAKE_MODEL = "gemini-3.5-flash-lite"  # fast/cheap model for profile parsing
GEMINI_REPORT_MODEL = "gemini-3.8-flash"       # larger model for the scenario report
GEMINI_FALLBACK_MODELS = ("gemini-3.5-flash", "gemini-flash-lite-latest")
RETRYABLE_STATUS_CODES = {429, 500, 502, 503, 504}
gemini_client = genai.Client(api_key=GEMINI_API_KEY)

# ==========================================
# STATE INITIALIZATIONS
# ==========================================
if "is_authenticated" not in st.session_state: st.session_state.is_authenticated = False
if "user_id" not in st.session_state: st.session_state.user_id = None
if "scenario_history" not in st.session_state: st.session_state.scenario_history = {} 

# Conversational UI States
if "flow_stage" not in st.session_state: st.session_state.flow_stage = "intake"
if "messages" not in st.session_state: st.session_state.messages = []
if "extracted_profile" not in st.session_state: st.session_state.extracted_profile = {}
if "report_data" not in st.session_state: st.session_state.report_data = None
if "hobby_assessment" not in st.session_state: st.session_state.hobby_assessment = None
if "salary_expectation" not in st.session_state: st.session_state.salary_expectation = None
if "active_history_key" not in st.session_state: st.session_state.active_history_key = None
if "pending_regenerate" not in st.session_state: st.session_state.pending_regenerate = False
# ==========================================
# AUTHENTICATION SYSTEM
# ==========================================
def hash_password(password):
    return hashlib.sha256(password.encode()).hexdigest()

def register_user(username, password, email):
    if not username or not username.strip():
        return False
    hashed_pw = hash_password(password)
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT user_id FROM users WHERE username = %s", (username,))
        if cursor.fetchone():
            return False
        cursor.execute("INSERT INTO users (username, password_hash, email) VALUES (%s, %s, %s)", (username, hashed_pw, email))
        conn.commit()
        return True
    except Exception:
        conn.rollback()
        raise
    finally:
        _close_quietly(cursor, conn)
        
def authenticate_user(username, password):
    hashed_pw = hash_password(password)
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT user_id, password_hash FROM users WHERE username = %s", (username,))
        result = cursor.fetchone()
    finally:
        _close_quietly(cursor, conn)
    if result and result[1] == hashed_pw:
        st.session_state.session_id = username
        st.session_state.user_id = result[0]
        st.session_state.is_authenticated = True
        hydrate_history_from_db()
        return True
    return False

def auth_or_register_google_user(google_name, email):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT user_id, username FROM users WHERE email = %s", (email,))
        existing = cursor.fetchone()
        if existing:
            user_id = existing[0]
            final_username = existing[1]
        else:
            base_username = google_name
            cursor.execute("SELECT user_id FROM users WHERE username = %s", (base_username,))
            if cursor.fetchone():
                base_username = f"{base_username}_{hashlib.sha256(email.encode()).hexdigest()[:4]}"
            cursor.execute("INSERT INTO users (username, password_hash, email) VALUES (%s, %s, %s)", (base_username, "GOOGLE_OAUTH", email))
            conn.commit()
            user_id = cursor.lastrowid
            final_username = base_username
    except Exception:
        conn.rollback()
        raise
    finally:
        _close_quietly(cursor, conn)
    st.session_state.session_id = final_username
    st.session_state.user_id = user_id
    st.session_state.is_authenticated = True
    hydrate_history_from_db()
    st.rerun()

def logout_user():
    for key in ["session_id", "user_id", "profile", "report_data", "hobby_assessment", "salary_expectation", "active_history_key"]:
        st.session_state[key] = None
    st.session_state.scenario_history = {}
    st.session_state.extracted_profile = {}
    st.session_state.messages = []
    st.session_state.flow_stage = "intake"
    st.session_state.is_authenticated = False
    st.rerun()

# ==========================================
# GENERATIVE AI ENGINE
# ==========================================
def _to_float(value, default=0.0):
    """Safely coerce a value returned by the LLM into a float."""
    try:
        return float(value)
    except (TypeError, ValueError):
        return default

def _extract_json(raw_text):
    """Strip markdown code fences and parse the model's JSON response."""
    if not raw_text:
        raise ValueError("The AI model returned an empty response.")
    cleaned = re.sub(r'```(?:json)?|```', '', raw_text).strip()
    match = re.search(r'\{.*\}', cleaned, re.DOTALL)
    if match:
        cleaned = match.group(0)
    return json.loads(cleaned)

def _generate_content(model, contents, fallback_models=GEMINI_FALLBACK_MODELS):
    """Call Gemini, retrying transient failures and falling back to other models.

    Preview/free-tier Gemini models regularly answer with ``503 UNAVAILABLE``
    ("high demand") or ``429`` rate-limit spikes, so a single call is not
    reliable enough for a live app.
    """
    last_error = None
    plan = [(model, 3)] + [(fb, 1) for fb in fallback_models if fb != model]
    for model_name, attempts in plan:
        for attempt in range(attempts):
            try:
                return gemini_client.models.generate_content(model=model_name, contents=contents)
            except genai_errors.APIError as exc:
                last_error = exc
                if getattr(exc, "code", None) not in RETRYABLE_STATUS_CODES:
                    raise
                time.sleep(2 * (attempt + 1))
    raise RuntimeError(f"Gemini request failed after retries: {last_error}") from last_error

# ---------------------------------------------------------------------------
# CAREER CASE MODEL
# Three supported cases: unemployed graduate / corporate shift / hobby->career
# ---------------------------------------------------------------------------
CASE_LABELS = {
    "unemployed_graduate": "Unemployed graduate / fresher",
    "corporate_shift": "Corporate / IT professional seeking a shift",
    "hobby_to_career": "Turning a hobby into a career",
}

CASE_DETAIL_FIELDS = {
    "unemployed_graduate": ["education", "graduation_year", "skills", "experience", "target_role"],
    "corporate_shift": ["current_company", "current_salary", "skills", "experience", "target_role"],
    "hobby_to_career": ["hobby", "target_role"],
}

FIELD_LABELS = {
    "education": "your degree and branch (e.g. B.Tech, Mechanical)",
    "graduation_year": "your pass-out year",
    "skills": "your skill set",
    "experience": "your hands-on experience",
    "current_company": "your current company",
    "current_salary": "your current salary",
    "hobby": "the hobby or interest you want to build a career on",
    "target_role": "the position you want to reach from here",
}

UNKNOWN_VALUES = {"", "-", "n/a", "na", "nil", "none", "null", "not specified",
                  "not sure", "unknown", "unavailable", "not mentioned", "not provided"}

def _is_unknown(value):
    """True when the model left a field blank or filled it with a placeholder."""
    if value is None:
        return True
    return str(value).strip().lower() in UNKNOWN_VALUES

def case_label(profile):
    return CASE_LABELS.get((profile or {}).get("case_type"), CASE_LABELS["unemployed_graduate"])

def case_key(profile):
    key = (profile or {}).get("case_type")
    return key if key in CASE_LABELS else "unemployed_graduate"

def missing_detail_fields(profile):
    """Fields the app still has to ask the user for, given the detected case."""
    return [f for f in CASE_DETAIL_FIELDS[case_key(profile)] if _is_unknown((profile or {}).get(f))]

def profile_summary(profile):
    """Human readable snapshot of everything captured so far."""
    profile = profile or {}
    rows = [
        ("Case", case_label(profile)),
        ("Current status", profile.get("current_job")),
        ("Education", profile.get("education")),
        ("Pass-out year", profile.get("graduation_year")),
        ("Current company", profile.get("current_company")),
        ("Current salary", profile.get("current_salary")),
        ("Skills", profile.get("skills")),
        ("Experience", profile.get("experience")),
        ("Hobby", profile.get("hobby")),
        ("Hobby evidence", profile.get("hobby_evidence")),
        ("Target position", profile.get("target_role")),
    ]
    return [f"**{name}:** {value}" for name, value in rows if not _is_unknown(value)]

# ---------------------------------------------------------------------------
# FILE / MEDIA HANDLING - attaching an artifact skips the typing
# ---------------------------------------------------------------------------
TEXT_TYPES = ["txt", "csv"]
DOC_TYPES = ["pdf", "docx"]
IMAGE_TYPES = ["png", "jpg", "jpeg", "webp"]
VIDEO_TYPES = ["mp4", "mov", "webm"]
UPLOAD_TYPES = TEXT_TYPES + DOC_TYPES + IMAGE_TYPES + VIDEO_TYPES

def wait_for_file_active(file_obj, timeout=180):
    """Video uploads are processed server side; wait until Gemini marks them ACTIVE."""
    deadline = time.time() + timeout
    while getattr(file_obj, "state", None) is not None:
        state = getattr(file_obj.state, "name", str(file_obj.state))
        if state != "PROCESSING":
            break
        if time.time() > deadline:
            raise TimeoutError("Gemini is still processing that video. Please try a shorter clip.")
        time.sleep(3)
        file_obj = gemini_client.files.get(name=file_obj.name)
    return file_obj

def load_uploaded_file(uploaded_file):
    """Convert a Streamlit UploadedFile into (text, gemini_media_part, label)."""
    if uploaded_file is None:
        return "", None, ""
    ext = uploaded_file.name.split(".")[-1].lower()
    if ext in TEXT_TYPES:
        return uploaded_file.getvalue().decode("utf-8", errors="ignore"), None, f"{ext.upper()} document"
    if ext in DOC_TYPES:
        if ext == "pdf":
            reader = PyPDF2.PdfReader(uploaded_file)
            text = " ".join(p.extract_text() for p in reader.pages if p.extract_text())
            label = "PDF document"
        else:
            document = docx.Document(uploaded_file)
            text = " ".join(p.text for p in document.paragraphs)
            label = "DOCX document"
        return text, None, label
    if ext in IMAGE_TYPES:
        mime = "image/jpeg" if ext in ("jpg", "jpeg") else f"image/{ext}"
        part = genai_types.Part.from_bytes(data=uploaded_file.getvalue(), mime_type=mime)
        return "[The user attached an image of their work.]", part, f"{ext.upper()} image"
    if ext in VIDEO_TYPES:
        with tempfile.NamedTemporaryFile(delete=False, suffix=f".{ext}") as tmp:
            tmp.write(uploaded_file.getvalue())
            path = tmp.name
        try:
            part = wait_for_file_active(gemini_client.files.upload(file=path))
        finally:
            try:
                os.remove(path)
            except OSError:
                pass
        return "[The user attached a workflow video.]", part, f"{ext.upper()} video"
    return "", None, f"{ext.upper()} file"

def build_contents(prompt, media_part=None):
    """Gemini takes [media, prompt] when media is attached, otherwise just the prompt."""
    return [media_part, prompt] if media_part else prompt
# ---------------------------------------------------------------------------
# AI INTAKE  (Case I, II and III all funnel through here)
# ---------------------------------------------------------------------------
INTAKE_SCHEMA = """
{
    "case_type": "unemployed_graduate | corporate_shift | hobby_to_career",
    "current_job": "their current role, or 'Unemployed Graduate', or 'Hobbyist'",
    "education": "degree and branch, e.g. 'B.Tech, Mechanical'",
    "graduation_year": "pass-out year, e.g. '2024'",
    "current_company": "current employer, only if they are employed",
    "current_salary": "current salary, only if they are employed",
    "skills": "comma separated list of skills that actually appear in the input",
    "experience": "e.g. '2 years IT', '6 months internship', 'Fresh Graduate'",
    "hobby": "the hobby or interest they want to turn into a career, if any",
    "hobby_evidence": "what they have actually produced or done in that hobby, if stated",
    "target_role": "the position they want to reach from this point",
    "notes": "anything else relevant"
}
"""

def ai_parse_intake(user_text, file_content="", gemini_media_file=None, existing_profile=None):
    known = json.dumps(existing_profile or {}, ensure_ascii=False)
    prompt = f"""
    You are the intake analyst for a career simulation app.
    Read the user's message, any attached resume / document text and any attached media
    (an image or a workflow video), then extract a structured professional profile.

    Profile already collected (never drop a value unless the new input corrects it):
    {known}

    New user message: {user_text}
    Resume / document text: {file_content}

    Pick the case_type that fits:
    - "unemployed_graduate": a student, fresher or unemployed person exploring a first career path.
    - "corporate_shift": currently employed (IT / corporate) and wants a role or company shift.
    - "hobby_to_career": wants to convert a hobby or personal interest into a career.

    Hard rules:
    - Report only facts that appear in the input. Use "" for anything not stated.
    - Never write placeholders such as "Not specified", "Unknown", "N/A" or "-". Use "" instead.
    - Never invent skills, employers, salaries or dates.
    - If the input is a resume, mine it thoroughly for education, skills and dates.

    Respond strictly with this JSON object and nothing else:
    {INTAKE_SCHEMA}
    """
    res = _generate_content(GEMINI_INTAKE_MODEL, build_contents(prompt, gemini_media_file))
    parsed = _extract_json(res.text)

    merged = dict(existing_profile or {})
    for key, value in parsed.items():
        if not _is_unknown(value):
            merged[key] = value
    if merged.get("case_type") not in CASE_LABELS:
        merged["case_type"] = case_key(existing_profile)
    return merged

# ---------------------------------------------------------------------------
# AI HOBBY PROFICIENCY  (Case III - analyse the work objectively)
# ---------------------------------------------------------------------------
PROFICIENCY_SCHEMA = """
{
    "proficiency": "Beginner | Intermediate | Advanced | Professional",
    "years_experience": "e.g. '2 years'",
    "tools": "comma separated tools / software they appear to work with",
    "strengths": ["strength 1", "strength 2"],
    "gaps": ["gap 1", "gap 2"],
    "portfolio_quality": "one sentence on the objective quality of the submitted work",
    "summary": "2-3 sentence objective assessment written directly to the user"
}
"""

def ai_assess_hobby_proficiency(hobby, user_text, file_content="", gemini_media_file=None, profile=None):
    context = json.dumps(profile or {}, ensure_ascii=False)
    prompt = f"""
    You are an expert assessor of practical, hands-on skill.
    Objectively judge how proficient this person really is at their hobby, using their own
    description and, when attached, the image / document / workflow video of their work.

    Hobby or interest: {hobby}
    Their own description: {user_text}
    Supporting document text: {file_content}
    Other known context about them: {context}

    If an image or a video is attached, examine the real work closely: technique, finish,
    complexity, tooling and the order of the workflow. Weigh that hard evidence more heavily
    than the person's self-description, and say so in the summary.

    Respond strictly with this JSON object and nothing else:
    {PROFICIENCY_SCHEMA}
    """
    res = _generate_content(GEMINI_INTAKE_MODEL, build_contents(prompt, gemini_media_file))
    return _extract_json(res.text)

REPORT_SCHEMA = """
{
    "scenarios": [
        {
            "name": "Scenario A: [Role Name]",
            "archetype": "Steady Specialist | Aggressive Growth | Adjacent Pivot | Independent / Freelance | Leadership Track | Global / Remote-First | Hybrid / Side-by-Side",
            "risk_level": "Low | Medium | High",
            "target_role": "[Specific Job Title]",
            "timeline_years": [Integer 1-10],
            "expected_salary": [Integer in INR],
            "expected_savings": [Integer in INR, assuming a 30% savings rate over the timeline],
            "present_vs_future": "One sentence comparing where they stand today with where this path lands them.",
            "benefits": ["Benefit 1", "Benefit 2", "Benefit 3"],
            "trade_offs": ["Trade-off 1", "Trade-off 2"]
        }
    ],
    "objective_verdict": "3 sentences comparing every scenario objectively, naming the strongest one and why.",
    "recommended_scenario": "[Exact scenario name whose benefits clearly outweigh its trade-offs, or 'No clear winner']",
    "decision_note": "A warm 2 sentence note that offers the recommendation as advice and leaves the final decision with the user."
}
"""

DEFAULT_SCENARIO_COUNT = 4
MIN_SCENARIO_COUNT = 2
MAX_SCENARIO_COUNT = 6

# How many years ahead the simulation projects. Kept within the 1-10 range the
# report schema allows for "timeline_years".
DEFAULT_SIMULATION_YEARS = 5
MIN_SIMULATION_YEARS = 1
MAX_SIMULATION_YEARS = 10

ARCHETYPES = [
    "Steady Specialist - go deep in the current skill set",
    "Aggressive Growth - a high-risk, high-reward role",
    "Adjacent Pivot - carry transferable skills into a neighbouring industry",
    "Independent / Freelance - build a personal client base",
    "Leadership Track - move toward managing people and budgets",
    "Global / Remote-First - serve overseas clients from India",
    "Hybrid / Side-by-Side - keep the current income while building the new path",
]

NO_SALARY_ANSWERS = {"no idea", "noidea", "no", "none", "n/a", "na", "idk", "i don't know",
                     "i dont know", "dont know", "don't know", "not sure", "unsure", "skip"}

# ==========================================
# PERSISTENCE LAYER (MySQL-backed history)
# ==========================================
# Every simulation is mirrored into MySQL so a refresh, a server restart or
# a second login restores it. scenario_history stays as the fast in-memory
# path; these helpers keep the DB copy in step with it.
PROFILE_COLUMNS = (
    "case_type", "current_job", "education", "graduation_year",
    "current_company", "current_salary", "skills", "experience",
    "hobby", "hobby_evidence", "target_role", "notes",
)

def _opt_int(value):
    """Coerce an INR figure to int; None when the model omitted it."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None

def _close_quietly(cursor, conn):
    try:
        cursor.close()
    except Exception:
        pass
    try:
        conn.close()
    except Exception:
        pass

def persist_simulation(user_id, profile, salary, report,
                       scenario_count=None, horizon_years=None):
    """Write one full simulation to MySQL; return its history title."""
    if not user_id or not isinstance(report, dict):
        return None
    scenarios = report.get("scenarios") or []
    try:
        scenario_count = int(scenario_count if scenario_count is not None else len(scenarios))
    except (TypeError, ValueError):
        scenario_count = len(scenarios) or DEFAULT_SCENARIO_COUNT
    if not scenario_count:
        scenario_count = len(scenarios) or DEFAULT_SCENARIO_COUNT
    scenario_count = max(MIN_SCENARIO_COUNT, min(MAX_SCENARIO_COUNT, scenario_count))
    try:
        horizon_years = int(horizon_years)
    except (TypeError, ValueError):
        horizon_years = DEFAULT_SIMULATION_YEARS
    horizon_years = max(MIN_SIMULATION_YEARS, min(MAX_SIMULATION_YEARS, horizon_years))
    profile = dict(profile or {})
    stamp = time.strftime("%d %b %Y, %H:%M")
    target = profile.get("target_role")
    title = stamp + (f" · {target}" if not _is_unknown(target) else "")
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("DELETE FROM reports WHERE user_id = %s AND run_key = %s",
                       (user_id, title))
        cursor.execute(
            "INSERT INTO profiles (user_id, case_type, current_job, education,"
            " graduation_year, current_company, current_salary, skills,"
            " experience, hobby, hobby_evidence, target_role, notes,"
            " salary_expectation)"
            " VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)",
            (user_id, profile.get("case_type"), profile.get("current_job"),
             profile.get("education"), profile.get("graduation_year"),
             profile.get("current_company"), profile.get("current_salary"),
             profile.get("skills"), profile.get("experience"),
             profile.get("hobby"), profile.get("hobby_evidence"),
             profile.get("target_role"), profile.get("notes"), salary),
        )
        profile_id = cursor.lastrowid
        cursor.execute(
            "INSERT INTO reports (user_id, profile_id, run_key, scenario_count,"
            " simulation_years, salary_expectation, objective_verdict,"
            " recommended_scenario, decision_note)"
            " VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)",
            (user_id, profile_id, title, scenario_count, horizon_years, salary,
             report.get("objective_verdict"),
             report.get("recommended_scenario"), report.get("decision_note")),
        )
        report_id = cursor.lastrowid
        for scenario in scenarios:
            if not isinstance(scenario, dict):
                continue
            try:
                years = int(scenario.get("timeline_years"))
            except (TypeError, ValueError):
                years = horizon_years
            years = max(MIN_SIMULATION_YEARS, min(MAX_SIMULATION_YEARS, years))
            cursor.execute(
                "INSERT INTO scenarios (report_id, user_id, run_key, name,"
                " archetype, risk_level, target_role, timeline_years,"
                " expected_salary, expected_savings, present_vs_future,"
                " benefits, trade_offs)"
                " VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s,"
                " CAST(%s AS JSON), CAST(%s AS JSON))",
                (report_id, user_id, title, scenario.get("name"),
                 scenario.get("archetype"), scenario.get("risk_level"),
                 scenario.get("target_role"), years,
                 _opt_int(scenario.get("expected_salary")),
                 _opt_int(scenario.get("expected_savings")),
                 scenario.get("present_vs_future"),
                 json.dumps(scenario.get("benefits") or []),
                 json.dumps(scenario.get("trade_offs") or [])),
            )
        conn.commit()
        return title
    except Exception:
        conn.rollback()
        raise
    finally:
        _close_quietly(cursor, conn)

def fetch_history_titles(user_id):
    """[(run_key, scenario_count, simulation_years)] newest first."""
    if not user_id:
        return []
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(
            "SELECT run_key, scenario_count, simulation_years FROM reports"
            " WHERE user_id = %s ORDER BY report_id DESC",
            (user_id,),
        )
        return list(cursor.fetchall())
    finally:
        _close_quietly(cursor, conn)

def fetch_simulation(user_id, run_key):
    """Rebuild one save_to_history() entry from MySQL (or None)."""
    if not user_id or not run_key:
        return None
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute(
            "SELECT r.scenario_count, r.simulation_years, r.salary_expectation,"
            " r.objective_verdict, r.recommended_scenario, r.decision_note,"
            " p.case_type, p.current_job, p.education, p.graduation_year,"
            " p.current_company, p.current_salary, p.skills, p.experience,"
            " p.hobby, p.hobby_evidence, p.target_role, p.notes"
            " FROM reports r LEFT JOIN profiles p ON p.profile_id = r.profile_id"
            " WHERE r.user_id = %s AND r.run_key = %s",
            (user_id, run_key),
        )
        head = cursor.fetchone()
        if not head:
            return None
        cursor.execute(
            "SELECT name, archetype, risk_level, target_role, timeline_years,"
            " expected_salary, expected_savings, present_vs_future, benefits,"
            " trade_offs FROM scenarios"
            " WHERE user_id = %s AND run_key = %s ORDER BY scenario_id",
            (user_id, run_key),
        )
        scenarios = []
        for row in cursor.fetchall():
            for list_key in ("benefits", "trade_offs"):
                value = row.get(list_key)
                if isinstance(value, str):
                    try:
                        value = json.loads(value)
                    except (TypeError, ValueError):
                        value = []
                row[list_key] = value or []
            scenarios.append(row)
        profile = {key: head.get(key) for key in PROFILE_COLUMNS}
        return {
            "report": {
                "scenarios": scenarios,
                "objective_verdict": head.get("objective_verdict"),
                "recommended_scenario": head.get("recommended_scenario"),
                "decision_note": head.get("decision_note"),
            },
            "profile": profile,
            "hobby_assessment": None,
            "salary": head.get("salary_expectation"),
            "scenario_count": head.get("scenario_count") or len(scenarios),
            "simulation_years": head.get("simulation_years") or DEFAULT_SIMULATION_YEARS,
        }
    finally:
        _close_quietly(cursor, conn)

def hydrate_history_from_db():
    """Reload this user's saved simulations into scenario_history on login."""
    user_id = st.session_state.get("user_id")
    history = {}
    if user_id:
        try:
            for run_key, _count, _years in fetch_history_titles(user_id):
                entry = fetch_simulation(user_id, run_key)
                if entry:
                    history[run_key] = entry
        except Exception as exc:
            st.warning(f"Saved simulations could not be loaded: {exc}")
    st.session_state.scenario_history = history
    st.session_state.active_history_key = None

def _apply_horizon(report, horizon_years):
    """Force every scenario onto the horizon chosen in Simulation Settings.

    The model is asked to use one shared timeline, but a stray value would put the
    side-by-side comparison and the wealth chart out of step with the setting, so
    the setting always wins.
    """
    if not isinstance(report, dict):
        return report
    for scenario in report.get("scenarios") or []:
        if isinstance(scenario, dict):
            scenario["timeline_years"] = horizon_years
    return report

def ai_generate_report(profile, salary_expectation, hobby_assessment=None,
                       scenario_count=DEFAULT_SCENARIO_COUNT,
                       horizon_years=DEFAULT_SIMULATION_YEARS):
    profile = profile or {}
    try:
        scenario_count = int(scenario_count)
    except (TypeError, ValueError):
        scenario_count = DEFAULT_SCENARIO_COUNT
    scenario_count = max(MIN_SCENARIO_COUNT, min(MAX_SCENARIO_COUNT, scenario_count))
    try:
        horizon_years = int(horizon_years)
    except (TypeError, ValueError):
        horizon_years = DEFAULT_SIMULATION_YEARS
    horizon_years = max(MIN_SIMULATION_YEARS, min(MAX_SIMULATION_YEARS, horizon_years))
    archetype_menu = "\n".join(
        f"      {index + 1}. {name}" for index, name in enumerate(ARCHETYPES[:scenario_count + 1])
    )
    wants_no_salary = _is_unknown(salary_expectation) or str(salary_expectation).strip().lower() in NO_SALARY_ANSWERS
    if wants_no_salary:
        salary_text = ("The user did not give a salary expectation. Do NOT demand one and do NOT invent "
                       "an aggressive number: proceed without it, using a realistic market range for the "
                       "role, city and experience in India (INR), and state that assumption briefly.")
    else:
        salary_text = str(salary_expectation)
    hobby_text = json.dumps(hobby_assessment, ensure_ascii=False) if hobby_assessment else "Not applicable."

    prompt = f"""
    You are an expert, honest career strategist. Build a branching scenario report.

    USER BASELINE (their present situation)
    - Case: {case_label(profile)}
    - Current status / job: {profile.get('current_job')}
    - Education: {profile.get('education')}
    - Pass-out year: {profile.get('graduation_year')}
    - Current company: {profile.get('current_company')}
    - Current salary: {profile.get('current_salary')}
    - Skills: {profile.get('skills')}
    - Experience: {profile.get('experience')}
    - Hobby or interest: {profile.get('hobby')}
    - Hobby evidence so far: {profile.get('hobby_evidence')}
    - Hobby proficiency assessment: {hobby_text}
    - Position the user wants to reach: {profile.get('target_role')}
    - Salary expectation: {salary_text}

    Create exactly {scenario_count} genuinely DIFFERENT future scenarios, each spanning a
    {horizon_years} year horizon.
    Give every scenario its own archetype, choosing the {scenario_count} that fit this user best from:
{archetype_menu}
    Every scenario must set timeline_years = {horizon_years} and estimate expected_savings over
    that same {horizon_years} year horizon, so the paths can be compared side by side.

    Each scenario must have a distinct target role, salary, risk level and trade-off
    profile. Do not simply rename the same plan {scenario_count} times - the user must be able to
    see real alternatives. Name them "Scenario A: ...", "Scenario B: ..." and so on, in order.

    For each one, compare the PRESENT situation with the FUTURE situation, and list its concrete
    benefits along with its honest trade-offs. Do not hide the downsides.

    Then give an objective verdict that compares ALL {scenario_count} scenarios and names the one
    whose benefits most clearly outweigh its trade-offs for THIS user. Only recommend when that
    advantage is genuine; if no path clearly wins, say "No clear winner".
    Finally write a short, warm note that offers the recommendation as advice and politely leaves
    the final decision with the user.

    Respond strictly with this JSON object and nothing else:
    {REPORT_SCHEMA}
    """
    res = _generate_content(GEMINI_REPORT_MODEL, prompt)
    return _apply_horizon(_extract_json(res.text), horizon_years)

# ==========================================
# DATA VISUALIZATION
# ==========================================
def create_financial_charts(scenarios_data):
    df_list = []
    for s in (scenarios_data or []):
        timeline_years = int(_to_float(s.get("timeline_years"), 0))
        expected_savings = _to_float(s.get("expected_savings"), 0.0)
        if timeline_years <= 0:
            continue
        annual_savings = expected_savings / timeline_years
        years = [0] + list(range(1, timeline_years + 1))
        trajectory = [0.0] + [annual_savings * y for y in range(1, timeline_years + 1)]
        temp_df = pd.DataFrame({"Year": years, "Accumulated_Savings": trajectory, "Scenario": s.get("name", "Scenario")})
        df_list.append(temp_df)

    if not df_list:
        return None

    combined_df = pd.concat(df_list)
    fig = px.line(combined_df, x="Year", y="Accumulated_Savings", color="Scenario", 
                  title="Projected Wealth Trajectory (Based on Target Savings)", markers=True)
    fig.update_traces(line=dict(width=3), marker=dict(size=8))
    return fig

# ==========================================
# CONVERSATIONAL UI HELPERS
# ==========================================
SUGGESTIONS = [
    ("🎓 I am an unemployed graduate looking to map out my career options",
     "I am an unemployed graduate looking to map out my career options."),
    ("🏢 I am a corporate IT employee looking to shift roles",
     "I am a corporate IT employee looking to shift roles."),
    ("🎨 I want to turn my hobby into a career",
     "I want to turn my hobby into a career."),
]

STAGE_PLACEHOLDERS = {
    "intake": "Ask FutureMe Simulator... (or attach your resume)",
    "details": "Type the missing details, or attach your resume...",
    "proficiency": "Describe your experience, or attach your workflow...",
    "salary": "Enter your salary expectation, or type 'No idea'...",
}

STAGE_HINTS = {
    "details": [("⏭ Skip these details", "skip")],
    "proficiency": [("⏭ Skip - assess from my description", "skip")],
    "salary": [("🤷 No idea - proceed without a salary target", "No idea")],
}

DEFAULT_DECISION_NOTE = (
    "These projections are simulations built from what you shared, not guarantees. "
    "The right path depends on what matters most to you, so please treat this as advice "
    "rather than an instruction. You can re-run the simulation whenever your situation changes."
)

def add_message(role, content):
    st.session_state.messages.append({"role": role, "content": content})

def render_hero():
    st.markdown(
        "<div class='fm-hero'>"
        "<h1>Where should we start?</h1>"
        "<p class='fm-hero-sub'>Tell me your current situation &mdash; or attach your resume, "
        "a photo of your work or a workflow video and I will read it for you.</p>"
        "</div>",
        unsafe_allow_html=True,
    )

def stage_question(stage, profile):
    """The single question the assistant asks for a given stage."""
    profile = profile or {}
    if stage == "details":
        missing = missing_detail_fields(profile)
        if not missing:
            return None
        bullets = "\n".join(f"- {FIELD_LABELS[f]}" for f in missing if f in FIELD_LABELS)
        return (f"I have logged this as a **{case_label(profile)}** case.\n\n"
                f"To simulate your future properly I still need:\n{bullets}\n\n"
                "You can send all of these in one message, **or attach your resume / a document** "
                "and I will extract them for you. Type **skip** if you would rather proceed with "
                "what we already have.")
    if stage == "proficiency":
        hobby = profile.get("hobby") or "your hobby"
        return (f"Before I simulate anything, let me gauge your level in **{hobby}** objectively.\n\n"
                "Tell me how long you have been doing it, which tools or software you use, and what "
                "you have produced so far. **Or attach a photo, a document or an MP4 video of your "
                "workflow** and I will assess the work itself. Type **skip** to let me judge from "
                "what we already know.")
    if stage == "salary":
        return ("One last thing before I run the simulation: **what are your exact salary "
                "expectations for this next move?**\n\n"
                "*Not sure? Just type **No idea** and I will run the simulation without a salary "
                "target.*")
    return None

def compute_next_stage():
    profile = st.session_state.extracted_profile or {}
    if missing_detail_fields(profile):
        return "details"
    if case_key(profile) == "hobby_to_career" and not st.session_state.hobby_assessment:
        return "proficiency"
    return "salary"

def advance(stage=None):
    """Move to a stage, append its question to the transcript and rerun."""
    if stage:
        st.session_state.flow_stage = stage
    question = stage_question(st.session_state.flow_stage, st.session_state.extracted_profile)
    if question and (not st.session_state.messages or st.session_state.messages[-1]["content"] != question):
        add_message("assistant", question)
    st.rerun()

def start_new_simulation():
    st.session_state.extracted_profile = {}
    st.session_state.report_data = None
    st.session_state.hobby_assessment = None
    st.session_state.salary_expectation = None
    st.session_state.messages = []
    st.session_state.active_history_key = None
    st.session_state.flow_stage = "intake"
    for key in [k for k in st.session_state.keys() if k.startswith("last_up_")]:
        st.session_state[key] = None
    st.rerun()

def load_history_entry(key):
    entry = st.session_state.scenario_history.get(key)
    if entry is None:
        entry = fetch_simulation(st.session_state.get("user_id"), key)
        if entry is None:
            st.error("That saved simulation could not be loaded.")
            return
        st.session_state.scenario_history[key] = entry
    st.session_state.report_data = entry["report"]
    st.session_state.extracted_profile = entry["profile"]
    st.session_state.hobby_assessment = entry.get("hobby_assessment")
    st.session_state.salary_expectation = entry.get("salary")
    st.session_state.scenario_count = entry.get("scenario_count", DEFAULT_SCENARIO_COUNT)
    st.session_state.simulation_years = entry.get("simulation_years", DEFAULT_SIMULATION_YEARS)
    st.session_state.active_history_key = key
    st.session_state.flow_stage = "report"
    st.rerun()

def save_to_history(report, profile, hobby_assessment, salary, scenario_count=None,
                    horizon_years=None):
    """History is recorded for the user so past simulations can be reopened.

    Writes to MySQL first (the durable copy) and mirrors the entry into the
    in-memory dict (the fast path). A DB failure no longer loses the user's
    data: the dict entry is still stored and a warning is shown instead.
    """
    profile = dict(profile or {})
    stamp = time.strftime("%d %b %Y, %H:%M")
    target = profile.get("target_role")
    fallback_title = stamp + (f" · {target}" if not _is_unknown(target) else "")
    title = fallback_title
    try:
        db_title = persist_simulation(
            st.session_state.get("user_id"), profile, salary, report,
            scenario_count, horizon_years,
        )
        if db_title:
            title = db_title
    except Exception as exc:
        st.warning(f"Simulation saved for this session, but the database copy failed: {exc}")
    st.session_state.scenario_history[title] = {
        "report": report,
        "profile": profile,
        "hobby_assessment": hobby_assessment,
        "salary": salary,
        "scenario_count": scenario_count or len(report.get("scenarios") or []),
        "simulation_years": horizon_years or DEFAULT_SIMULATION_YEARS,
    }
    st.session_state.active_history_key = title

def prompt_bar(stage, allow_attach=True):
    """The prompt bar: type-and-send plus an attach button that skips the questions."""
    uploaded_file = None
    user_text = ""
    submitted = False
    with st.container(border=True):
        attach_col, form_col = st.columns([1, 11], vertical_alignment="center")
        with attach_col:
            if allow_attach:
                with st.popover("📎", help="Attach a resume, certificate, photo or workflow video"):
                    uploaded_file = st.file_uploader(
                        "Attach a file", type=UPLOAD_TYPES, key=f"up_{stage}",
                        label_visibility="collapsed",
                    )
        with form_col:
            with st.form(key=f"form_{stage}", clear_on_submit=True, border=False):
                text_col, send_col = st.columns([12, 1], vertical_alignment="center")
                with text_col:
                    user_text = st.text_input(
                        "prompt", key=f"txt_{stage}",
                        label_visibility="collapsed",
                        placeholder=STAGE_PLACEHOLDERS[stage],
                    )
                with send_col:
                    submitted = st.form_submit_button("➤", type="primary")
    return user_text, uploaded_file, submitted

def newly_attached(stage, uploaded_file):
    """True exactly once per newly attached file - attaching alone can skip the questions."""
    if uploaded_file is None:
        return False
    signature = f"{uploaded_file.name}:{getattr(uploaded_file, 'size', 0)}"
    key = f"last_up_{stage}"
    if st.session_state.get(key) == signature:
        return False
    st.session_state[key] = signature
    return True
def render_prompt_stage():
    """Prompt bar first, then the suggestions / shortcuts underneath it."""
    stage = st.session_state.flow_stage
    text, uploaded_file, submitted = prompt_bar(stage)

    suggestion = None
    if stage == "intake":
        st.markdown("💡 **Suggestions from the application**")
        for index, (label, seed) in enumerate(SUGGESTIONS):
            if st.button(label, key=f"sug_{index}", width="stretch"):
                suggestion = seed
    else:
        for index, (label, payload) in enumerate(STAGE_HINTS.get(stage, [])):
            if st.button(label, key=f"hint_{stage}_{index}", width="stretch"):
                suggestion = payload

    attached = newly_attached(stage, uploaded_file)

    if suggestion:
        handle_turn(suggestion, uploaded_file)
    elif submitted and text.strip():
        handle_turn(text, uploaded_file)
    elif attached:
        handle_turn("", uploaded_file)

def handle_turn(text, uploaded_file):
    """Process one conversational turn for whichever stage is active."""
    stage = st.session_state.flow_stage
    profile = st.session_state.extracted_profile or {}
    text = (text or "").strip()

    file_text, media_part, label = "", None, ""
    if uploaded_file is not None:
        with st.spinner("Reading your attachment..."):
            try:
                file_text, media_part, label = load_uploaded_file(uploaded_file)
            except Exception as exc:
                st.error(f"I could not read that file: {exc}")
                return

    add_message("user", text or f"📎 [{label or 'attachment'}]")

    with st.chat_message("assistant"):
        if stage in ("intake", "details"):
            with st.spinner("Analysing your profile..."):
                try:
                    profile = ai_parse_intake(text, file_text, media_part, profile)
                except Exception as exc:
                    st.error(f"I could not read that: {exc}")
                    return
            if text.lower() in {"skip", "skip.", "skip it"}:
                for field in missing_detail_fields(profile):
                    profile[field] = "Not provided by user"
            st.session_state.extracted_profile = profile
            st.markdown("Here is what I have captured so far:")
            st.markdown("\n\n".join(profile_summary(profile)) or "_Nothing captured yet._")

            # Case III: if they already handed us a workflow artifact, assess it right now
            # instead of asking them to upload it again.
            if (case_key(profile) == "hobby_to_career" and not st.session_state.hobby_assessment
                    and (media_part is not None or file_text.strip())):
                with st.spinner("Assessing your work objectively..."):
                    try:
                        st.session_state.hobby_assessment = ai_assess_hobby_proficiency(
                            profile.get("hobby") or "the hobby in question",
                            text, file_text, media_part, profile,
                        )
                    except Exception:
                        pass

            next_stage = compute_next_stage()
            question = stage_question(next_stage, profile)
            if question:
                st.write(question)
                add_message("assistant", question)
            st.session_state.flow_stage = next_stage
            st.rerun()

        elif stage == "proficiency":
            if text:
                description = text
            elif media_part is not None or file_text.strip():
                description = ("The user supplied this work as the primary evidence and did not add "
                               "a written description.")
            else:
                description = ("The user skipped the questions. Judge from the profile only and say "
                               "clearly that you are doing so.")
            with st.spinner("Assessing your work objectively..."):
                try:
                    assessment = ai_assess_hobby_proficiency(
                        profile.get("hobby") or "the hobby in question",
                        description, file_text, media_part, profile,
                    )
                except Exception as exc:
                    st.error(f"I could not assess that: {exc}")
                    return
            st.session_state.hobby_assessment = assessment
            summary = (f"**Proficiency:** {assessment.get('proficiency')} · "
                       f"**Experience:** {assessment.get('years_experience')}\n\n"
                       f"{assessment.get('summary')}")
            st.markdown(summary)
            add_message("assistant", summary)
            advance("salary")

        elif stage == "salary":
            st.session_state.salary_expectation = None if text.lower() in NO_SALARY_ANSWERS else text
            scenario_count = st.session_state.get("scenario_count", DEFAULT_SCENARIO_COUNT)
            horizon_years = st.session_state.get("simulation_years", DEFAULT_SIMULATION_YEARS)
            with st.spinner(f"Simulating {scenario_count} branching futures over "
                            f"{horizon_years} years..."):
                try:
                    report = ai_generate_report(profile, text, st.session_state.hobby_assessment,
                                                scenario_count, horizon_years)
                except Exception as exc:
                    st.error(f"I could not generate the report: {exc}")
                    return
            st.session_state.report_data = report
            message = "Simulation complete! Here is your branching scenario report."
            st.write(message)
            add_message("assistant", message)
            save_to_history(report, profile, st.session_state.hobby_assessment, text,
                            scenario_count, horizon_years)
            st.session_state.flow_stage = "report"
            st.rerun()

def request_regenerate():
    """Callback: only flip a flag, so widget-keyed state is never touched here."""
    st.session_state.pending_regenerate = True

def run_regeneration():
    """Re-run the simulation with the scenario count and years set in the sidebar.

    Called from the script body (not from a callback) so the spinner and any
    error message render properly.
    """
    st.session_state.pending_regenerate = False
    scenario_count = st.session_state.get("scenario_count", DEFAULT_SCENARIO_COUNT)
    horizon_years = st.session_state.get("simulation_years", DEFAULT_SIMULATION_YEARS)
    profile = st.session_state.extracted_profile or {}
    with st.spinner(f"Re-simulating {scenario_count} branching futures over "
                    f"{horizon_years} years..."):
        try:
            report = ai_generate_report(profile, st.session_state.salary_expectation,
                                        st.session_state.hobby_assessment, scenario_count,
                                        horizon_years)
        except Exception as exc:
            st.error(f"I could not regenerate the report: {exc}")
            return
    st.session_state.report_data = report
    save_to_history(report, profile, st.session_state.hobby_assessment,
                    st.session_state.salary_expectation, scenario_count, horizon_years)
    st.rerun()

def render_report():
    if st.session_state.get("pending_regenerate"):
        run_regeneration()
    report = st.session_state.report_data or {}
    profile = st.session_state.extracted_profile or {}
    assessment = st.session_state.hobby_assessment
    scenarios = report.get("scenarios") or []

    st.markdown("---")
    st.markdown("## 📊 Branching Scenario Report")
    if st.session_state.active_history_key:
        st.caption(f"Saved to history · {st.session_state.active_history_key}")

    st.markdown("### Section A · Your Current Baseline")
    baseline = profile_summary(profile)
    st.info("\n\n".join(baseline) if baseline else "No baseline was captured.")
    if _is_unknown(st.session_state.salary_expectation):
        st.caption("Salary expectation: none supplied — projections use market estimates.")
    else:
        st.caption(f"Salary expectation used: {st.session_state.salary_expectation}")

    if assessment:
        with st.expander("🎯 Objective hobby proficiency assessment", expanded=True):
            st.markdown(f"**Proficiency:** {assessment.get('proficiency')} · "
                        f"**Experience:** {assessment.get('years_experience')}")
            if not _is_unknown(assessment.get("tools")):
                st.markdown(f"**Tools:** {assessment.get('tools')}")
            if not _is_unknown(assessment.get("portfolio_quality")):
                st.markdown(f"**Quality of the work you submitted:** {assessment.get('portfolio_quality')}")
            for strength in assessment.get("strengths") or []:
                st.markdown(f"- ✅ {strength}")
            for gap in assessment.get("gaps") or []:
                st.markdown(f"- 🔺 {gap}")
            st.write(assessment.get("summary"))

    horizon_years = st.session_state.get("simulation_years", DEFAULT_SIMULATION_YEARS)
    st.markdown(f"### Section B · Predicted Future Scenarios ({len(scenarios)} paths)")
    st.caption(f"Horizon: {horizon_years} years · {len(scenarios)} paths compared")
    if scenarios:
        columns = st.columns(2)
        for index, scenario in enumerate(scenarios):
            with columns[index % 2]:
                with st.container(border=True):
                    st.markdown(f"#### {scenario.get('name')}")
                    badges = []
                    if not _is_unknown(scenario.get("archetype")):
                        badges.append(f"🏷️ {scenario.get('archetype')}")
                    if not _is_unknown(scenario.get("risk_level")):
                        badges.append(f"⚡ Risk: {scenario.get('risk_level')}")
                    if badges:
                        st.caption(" · ".join(badges))
                    st.markdown(f"**Target Role:** {scenario.get('target_role')}")
                    st.markdown(f"**Timeline:** {scenario.get('timeline_years')} years")
                    st.markdown(f"**Expected Salary:** ₹{_to_float(scenario.get('expected_salary')):,.2f}")
                    st.markdown(f"**Projected Savings:** ₹{_to_float(scenario.get('expected_savings')):,.2f}")
                    benefits = scenario.get("benefits") or []
                    trade_offs = scenario.get("trade_offs") or []
                    benefit_metric, trade_metric = st.columns(2)
                    benefit_metric.metric("Benefits", len(benefits))
                    trade_metric.metric("Trade-offs", len(trade_offs))
                    if not _is_unknown(scenario.get("present_vs_future")):
                        st.caption(scenario["present_vs_future"])
                    st.markdown("✅ **Benefits**")
                    for benefit in benefits:
                        st.markdown(f"- {benefit}")
                    st.markdown("⚠️ **Trade-offs**")
                    for trade_off in trade_offs:
                        st.markdown(f"- {trade_off}")
    else:
        st.info("No scenarios were returned.")

    st.markdown("### Section C · Objective Verdict")
    recommended = report.get("recommended_scenario")
    if not _is_unknown(recommended) and str(recommended).strip().lower() not in {"no clear winner", "none"}:
        st.success(f"🤖 **Benefits clearly outweigh the trade-offs in: {recommended}**")
    else:
        st.warning("🤖 These paths are genuinely balanced — no scenario clearly outweighs the others.")
    st.write(report.get("objective_verdict") or "No verdict was returned.")

    if scenarios:
        st.markdown("#### Side-by-side comparison")
        comparison = pd.DataFrame([
            {
                "Scenario": scenario.get("name") or f"Path {index + 1}",
                "Archetype": scenario.get("archetype") or "-",
                "Target role": scenario.get("target_role") or "-",
                "Years": _to_float(scenario.get("timeline_years")),
                "Salary (₹)": _to_float(scenario.get("expected_salary")),
                "Savings (₹)": _to_float(scenario.get("expected_savings")),
                "Risk": scenario.get("risk_level") or "-",
                "Benefits": len(scenario.get("benefits") or []),
                "Trade-offs": len(scenario.get("trade_offs") or []),
                "Net (benefits - trade-offs)": (len(scenario.get("benefits") or [])
                                                - len(scenario.get("trade_offs") or [])),
            }
            for index, scenario in enumerate(scenarios)
        ])
        st.dataframe(comparison, width="stretch", hide_index=True)

    trajectory_fig = create_financial_charts(scenarios)
    if trajectory_fig is not None:
        st.plotly_chart(trajectory_fig, width="stretch")
    else:
        st.info("Not enough scenario data to plot a wealth trajectory.")

    st.markdown("### Section D · Over to You")
    st.info(report.get("decision_note") or DEFAULT_DECISION_NOTE)

    st.markdown("---")
    action_a, action_b, action_c = st.columns(3)
    action_a.button("🔄 Regenerate report", on_click=request_regenerate, width="stretch",
                    type="primary",
                    help="Re-run the simulation using the scenarios and years set in the sidebar")
    action_b.button("🔁 Simulate another future", on_click=start_new_simulation, width="stretch")
    action_c.button("🚪 Logout", on_click=logout_user, width="stretch")

# ==========================================
# STREAMLIT UI INTEGRATION
# ==========================================
st.set_page_config(page_title="FutureMe Simulator", layout="wide", initial_sidebar_state="expanded")

if not st.session_state.is_authenticated:
    if "code" in st.query_params:
        code = st.query_params["code"]
        st.query_params.clear() 
        token_url = "https://oauth2.googleapis.com/token"
        data = {"code": code, "client_id": GOOGLE_CLIENT_ID, "client_secret": GOOGLE_CLIENT_SECRET, "redirect_uri": REDIRECT_URI, "grant_type": "authorization_code"}
        res = requests.post(token_url, data=data)
        if res.status_code == 200:
            access_token = res.json().get("access_token")
            user_res = requests.get("https://www.googleapis.com/oauth2/v1/userinfo", headers={"Authorization": f"Bearer {access_token}"})
            if user_res.status_code == 200:
                user_data = user_res.json()
                auth_or_register_google_user(user_data.get("name", "GoogleUser"), user_data.get("email"))
    
    if os.path.exists("style.css"):
        with open("style.css") as f:
            st.markdown(f"<style>{f.read()}</style>", unsafe_allow_html=True)
            
    st.markdown("""
    <div class="left-panel">
        <div class="hero-logo">✳</div>
        <div class="hero-text">Hello<br>FutureMe! 👋</div>
        <div class="hero-sub">Skip repetitive and manual planning tasks. Get highly productive through simulation and save tons of time!</div>
        <div class="footer-text">© 2026 FutureMe Simulator. All rights reserved.</div>
    </div>
    """, unsafe_allow_html=True)
            
    st.markdown("<div class='brand-title'>FutureMe Simulator</div>", unsafe_allow_html=True)
    auth_tab1, auth_tab2 = st.tabs(["Login", "Register"])
    auth_url = f"https://accounts.google.com/o/oauth2/v2/auth?response_type=code&client_id={GOOGLE_CLIENT_ID}&redirect_uri={REDIRECT_URI}&scope=email%20profile"
    
    google_button_html = f"""
    <a href="{auth_url}" target="_self" style="display: block; width: 100%; text-align: center; background-color: #FFFFFF; color: #000000; border: 1px solid #E0E0E0; border-radius: 6px; padding: 1.4rem 0; font-weight: 600; margin-top: 0.5rem; text-decoration: none; font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif; transition: background-color 0.2s;">
        <img src="https://upload.wikimedia.org/wikipedia/commons/c/c1/Google_%22G%22_logo.svg" style="width: 18px; height: 18px; vertical-align: middle; margin-right: 8px; margin-bottom: 2px;">
        Login with Google
    </a>
    """
    
    with auth_tab1:
        st.markdown("<h1 style='font-size: 2.5rem; font-weight: 700; margin-bottom: 0.5rem;'>Welcome Back!</h1>", unsafe_allow_html=True)
        st.markdown("<p style='color: #666; font-size: 0.9rem; margin-bottom: 2rem;'>Don't have an account? <b style='color:black;'>Create a new account now</b>, it's FREE! Takes less than a minute.</p>", unsafe_allow_html=True)
        log_user = st.text_input("Username", key="log_user", label_visibility="collapsed", placeholder="Username")
        log_pass = st.text_input("Password", type="password", key="log_pass", label_visibility="collapsed", placeholder="Password")
        if st.button("Login Now", type="primary", width="stretch"):
            if authenticate_user(log_user, log_pass): st.rerun()
            else: st.error("Invalid credentials.")
        st.markdown(google_button_html, unsafe_allow_html=True)
            
    with auth_tab2:
        st.markdown("<h1 style='font-size: 2.5rem; font-weight: 700; margin-bottom: 0.5rem;'>Create Account</h1>", unsafe_allow_html=True)
        st.markdown("<p style='color: #666; font-size: 0.9rem; margin-bottom: 2rem;'>Already have an account? Switch to the <b style='color:black;'>Login tab</b> to access your vault.</p>", unsafe_allow_html=True)
        reg_user = st.text_input("New Username", key="reg_user", label_visibility="collapsed", placeholder="New Username")
        reg_email = st.text_input("Email", key="reg_email", label_visibility="collapsed", placeholder="Email Address")
        reg_pass = st.text_input("New Password", type="password", key="reg_pass", label_visibility="collapsed", placeholder="Password")
        if st.button("Register Now", type="primary", width="stretch"):
            if register_user(reg_user, reg_pass, reg_email): st.success("Registered successfully! Please switch to Login.")
            else: st.error("Username already exists.")
        st.markdown(google_button_html, unsafe_allow_html=True)

else:
    st.markdown("""
    <style>
    .left-panel { display: none !important; }
    .block-container { margin: 0 auto !important; max-width: 1000px !important; padding: 3rem 1rem !important; background: transparent !important; }
    .stApp { background-color: #0e1117 !important; }
    h1, h2, h3, p, label, span, div { color: #FAFAFA; }
    div[data-testid="stChatMessage"] { background-color: #1E1E24; border-radius: 10px; padding: 1rem; margin-bottom: 1rem; border: 1px solid #2D2D35; }
    div[data-testid="stChatMessage"][data-baseweb="card"] { background-color: transparent; border: none; }
    .fm-hero { text-align: center; padding: 3.5rem 0 0.75rem 0; }
    .fm-hero h1 { font-size: 2.75rem; font-weight: 800; margin: 0 0 0.6rem 0; letter-spacing: -0.5px; }
    .fm-hero-sub { font-size: 1rem; color: #9AA0A6 !important; margin: 0 auto 0.5rem auto; max-width: 640px; }
    div[data-testid="stForm"] { padding: 0 !important; }
    </style>
    """, unsafe_allow_html=True)
    
    # ==================================================================
    # SIDEBAR - history is recorded for the user
    # ==================================================================
    with st.sidebar:
        st.title("🎛️ Simulation Settings")
        if "scenario_count" not in st.session_state:
            st.session_state.scenario_count = DEFAULT_SCENARIO_COUNT
        if "simulation_years" not in st.session_state:
            st.session_state.simulation_years = DEFAULT_SIMULATION_YEARS
        st.slider(
            "Future scenarios to compare",
            min_value=MIN_SCENARIO_COUNT,
            max_value=MAX_SCENARIO_COUNT,
            key="scenario_count",
            help="How many distinct future paths the report should generate. "
                 "Change it, then press Regenerate report to re-run.",
        )
        st.slider(
            "No. of years",
            min_value=MIN_SIMULATION_YEARS,
            max_value=MAX_SIMULATION_YEARS,
            key="simulation_years",
            help="How many years ahead every scenario should project. Each path is "
                 "reported over this timeline, then press Regenerate report to re-run.",
        )
        st.markdown("---")
        st.title("📚 History")
        st.markdown("Every simulation you run is saved here:")
        if st.session_state.scenario_history:
            for key in list(st.session_state.scenario_history.keys()):
                prefix = "✅" if key == st.session_state.active_history_key else "📄"
                st.button(f"{prefix} {key}", key=f"hist_{key}", width="stretch",
                          on_click=load_history_entry, args=(key,),
                          help="Reopen this saved simulation")
        else:
            st.info("No history yet.")
        st.markdown("---")
        st.button("Logout", on_click=logout_user, width="stretch")
        st.button("Reset Conversation", on_click=start_new_simulation, width="stretch")

    # ==================================================================
    # TRANSCRIPT
    # ==================================================================
    for msg in st.session_state.messages:
        with st.chat_message(msg["role"]):
            st.write(msg["content"])

    # ==================================================================
    # STAGE ROUTER
    # ==================================================================
    if st.session_state.flow_stage == "report":
        render_report()
    else:
        if st.session_state.flow_stage == "intake" and not st.session_state.messages:
            render_hero()
        render_prompt_stage()
